from app.services.medical import extract_candidates


def test_candidate_has_traceability():
    candidates = extract_candidates([{
        "page": 2,
        "text": "Hemoglobin: 13.2 g/dl\nGlucose: 94 mg/dl",
        "method": "pdf_text",
    }])
    assert len(candidates) == 2
    assert candidates[0]["field"] == "hemoglobin"
    assert candidates[0]["source_page"] == 2
    assert candidates[0]["source_excerpt"] == "Hemoglobin: 13.2 g/dl"
    assert candidates[0]["status"] == "requires_human_review"


def test_candidate_is_not_a_diagnosis():
    candidates = extract_candidates([{"page": 1, "text": "HbA1c: 6.1 %", "method": "tesseract_ocr"}])
    assert candidates[0]["field"] == "hba1c"
    assert candidates[0]["confidence"] < 1.0
    assert candidates[0]["status"] == "requires_human_review"


def test_docx_extracts_paragraphs_and_tables(tmp_path):
    from docx import Document
    from app.services.medical import extract_document

    path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("Patient report")
    document.add_paragraph("Hemoglobin: 12.5 g/dL")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Glucose"
    table.cell(0, 1).text = "95 mg/dL"
    document.save(path)

    result = extract_document(str(path), use_cloud_ocr=False)
    assert result["extraction_method"] in {"docx_text", "pdf_text+tesseract_ocr", "mixed"}
    assert result.get("document_conversion") in {None, "libreoffice_to_pdf"}
    assert "Hemoglobin: 12.5 g/dL" in result["text"]
    assert ("Glucose: 95 mg/dL" in result["text"] or "Glucose 95 mg/dL" in result["text"] or "Glucose\n95 mg/dL" in result["text"])
    candidates = extract_candidates(result["pages"])
    assert {item["field"] for item in candidates} == {"hemoglobin", "glucose"}


def test_pdf_text_layer_detects_private_use_font_mapping():
    from app.services.medical import _text_layer_has_misencoded_glyphs

    assert _text_layer_has_misencoded_glyphs("\uf052\uf065\uf070\uf06f\uf072\uf074 Name")
    assert not _text_layer_has_misencoded_glyphs("Name: SAMPLE  Lab no: ABC123  Result: 12.3 mg/dL")


def test_machine_readable_payloads_append_to_ocr_without_losing_text(monkeypatch):
    import sys
    import types
    from types import SimpleNamespace
    from app.services.medical import _append_machine_readable_codes

    package = types.ModuleType("pyzbar")
    package.__path__ = []
    decoder = types.ModuleType("pyzbar.pyzbar")
    decoder.decode = lambda image: [SimpleNamespace(data=b"https://example.test/vaccine/123")]
    monkeypatch.setitem(sys.modules, "pyzbar", package)
    monkeypatch.setitem(sys.modules, "pyzbar.pyzbar", decoder)

    result = _append_machine_readable_codes("Vaccine certificate", object())
    assert result.startswith("Vaccine certificate")
    assert "Decoded QR/barcode payload" in result
    assert "https://example.test/vaccine/123" in result
